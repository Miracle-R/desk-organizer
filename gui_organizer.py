import sys
import os
import shutil
import csv
import yaml
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import messagebox, filedialog, scrolledtext

import docx
import pdfplumber
from pptx import Presentation
import openpyxl
import magic


# ---------- 帮助函数：获取打包后的资源路径 ----------
def resource_path(relative_path):
    """ 获取打包后或开发环境中文件的绝对路径 """
    try:
        # PyInstaller 会创建临时文件夹 _MEIPASS
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


# ---------- 文本提取与分类逻辑（与之前相同）----------
def load_categories():
    config_path = resource_path('categories.yaml')
    with open(config_path, 'r', encoding='utf-8') as f:
        data = yaml.safe_load(f)
    return data['categories']


def extract_text(file_path):
    suffix = file_path.suffix.lower()
    try:
        mime = magic.Magic(mime=True).from_file(str(file_path))
    except Exception:
        mime = ''
    try:
        if suffix in ('.docx', '.wps') or 'wordprocessingml' in mime:
            doc = docx.Document(file_path)
            return '\n'.join([p.text for p in doc.paragraphs])
        if suffix in ('.xlsx', '.et') or 'spreadsheetml' in mime:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            texts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_text = ' '.join([str(c) for c in row if c is not None])
                    if row_text.strip():
                        texts.append(row_text)
            return '\n'.join(texts)
        if suffix in ('.pptx', '.dps') or 'presentationml' in mime:
            prs = Presentation(file_path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        texts.append(shape.text)
            return '\n'.join(texts)
        if suffix == '.pdf' or 'pdf' in mime:
            text = ''
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages[:3]:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + '\n'
            return text
        if suffix in ('.txt', '.md'):
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        if suffix == '.csv':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                rows = [' '.join(row) for row in reader if any(row)]
                return '\n'.join(rows)
        return ''
    except Exception:
        return ''


def classify_document(text, categories):
    title_region = text[:200].lower()
    full_text = text.lower()
    scores = {}
    for cat in categories:
        score = 0
        tit_kws = cat.get('title_keywords', [])
        full_kws = cat.get('fulltext_keywords', [])
        for kw in tit_kws:
            score += title_region.count(kw) * 3
        for kw in full_kws:
            score += full_text.count(kw)
        if score > 0:
            scores[cat['name']] = score
    if not scores:
        return '未分类'
    return max(scores, key=scores.get)


def safe_move(src, dst_dir, log_file=None):
    dst_dir.mkdir(parents=True, exist_ok=True)
    dst = dst_dir / src.name
    base = src.stem
    ext = src.suffix
    counter = 1
    while dst.exists():
        dst = dst_dir / f'{base}_{counter}{ext}'
        counter += 1
    shutil.move(str(src), str(dst))
    if log_file:
        with open(log_file, 'a', encoding='utf-8') as f:
            f.write(f'{src}|{dst}\n')
    return dst


# ---------- GUI 应用类 ----------
class DeskOrganizerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("智能桌面整理器 v0.3")
        self.root.geometry("700x550")
        self.root.resizable(True, True)

        self.categories = load_categories()
        self.desktop = Path.home() / 'Desktop'

        # 文件列表显示区域
        frame_top = tk.Frame(root)
        frame_top.pack(pady=5)
        tk.Label(frame_top, text="桌面上的文件：", font=('微软雅黑', 10, 'bold')).pack(anchor='w')
        self.file_listbox = tk.Listbox(frame_top, width=90, height=15, selectmode=tk.EXTENDED)
        self.file_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = tk.Scrollbar(frame_top, orient=tk.VERTICAL, command=self.file_listbox.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.file_listbox.config(yscrollcommand=scrollbar.set)

        # 按钮区域
        btn_frame = tk.Frame(root)
        btn_frame.pack(pady=10)
        tk.Button(btn_frame, text="刷新列表", width=15, command=self.refresh_list).grid(row=0, column=0, padx=5)
        tk.Button(btn_frame, text="模拟预览", width=15, command=self.dry_run).grid(row=0, column=1, padx=5)
        tk.Button(btn_frame, text="开始整理", width=15, command=self.start_organize).grid(row=0, column=2, padx=5)
        tk.Button(btn_frame, text="一键撤回", width=15, command=self.undo_organize).grid(row=0, column=3, padx=5)

        # 状态输出区域
        self.status_text = scrolledtext.ScrolledText(root, width=90, height=10, state='normal')
        self.status_text.pack(pady=5, fill=tk.BOTH, expand=True)

        self.refresh_list()

    def log(self, msg):
        self.status_text.insert(tk.END, msg + '\n')
        self.status_text.see(tk.END)
        self.root.update()

    def refresh_list(self):
        self.file_listbox.delete(0, tk.END)
        files = [f for f in self.desktop.iterdir() if f.is_file() and not f.name.startswith('.') and f.suffix.lower() not in ('.lnk', '.ini')]
        for f in files:
            self.file_listbox.insert(tk.END, f.name)
        self.log(f'已扫描，发现 {len(files)} 个文件。')

    def get_files_from_listbox(self):
        selected_indices = self.file_listbox.curselection()
        all_files = [f for f in self.desktop.iterdir() if f.is_file() and not f.name.startswith('.') and f.suffix.lower() not in ('.lnk', '.ini')]
        if selected_indices:
            return [all_files[i] for i in selected_indices]
        else:
            return all_files

    def dry_run(self):
        files = self.get_files_from_listbox()
        if not files:
            messagebox.showinfo("提示", "没有文件可预览。")
            return
        self.log("===== 模拟预览 =====")
        categories = self.categories
        for file in files:
            text = extract_text(file)
            if not text.strip():
                cat = '未分类（无法读取内容）'
            else:
                cat = classify_document(text, categories)
            self.log(f'{file.name} -> {cat}')
        self.log("预览结束，未移动任何文件。\n")

    def start_organize(self):
        files = self.get_files_from_listbox()
        if not files:
            messagebox.showinfo("提示", "没有文件可整理。")
            return
        if not messagebox.askyesno("确认整理", f"即将整理 {len(files)} 个文件，是否继续？\n\n整理后的文件夹将创建在桌面上。"):
            return
        archive_root = self.desktop / f'桌面整理_{datetime.now().strftime("%Y%m%d_%H%M%S")}'
        log_path = archive_root / 'undo_log.txt'
        archive_root.mkdir(parents=True, exist_ok=True)

        self.log(f"开始整理到：{archive_root}")
        categories = self.categories
        for file in files:
            text = extract_text(file)
            if not text.strip():
                cat = '未分类（无法读取内容）'
            else:
                cat = classify_document(text, categories)
            try:
                dest = safe_move(file, archive_root / cat, log_file=log_path)
                self.log(f'✅ {file.name} -> {cat}')
            except Exception as e:
                self.log(f'❌ 移动失败 {file.name}: {e}')
        self.log(f"整理完成！撤回日志：{log_path}\n")
        messagebox.showinfo("完成", f"整理完成！\n如需撤回，请点击“一键撤回”并选择日志文件：\n{log_path}")
        self.refresh_list()

    def undo_organize(self):
        log_path = filedialog.askopenfilename(title="选择撤回日志文件", filetypes=[("日志文件", "undo_log.txt"), ("所有文件", "*.*")])
        if not log_path:
            return
        log_path = Path(log_path)
        if not log_path.exists():
            messagebox.showerror("错误", "找不到日志文件。")
            return
        if not messagebox.askyesno("确认撤回", f"将根据日志恢复文件，是否继续？\n{log_path}"):
            return
        archive_root = log_path.parent
        with open(log_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        for line in lines:
            src, dst = line.strip().split('|')
            src, dst = Path(src), Path(dst)
            if dst.exists():
                shutil.move(str(dst), str(src))
                self.log(f'已恢复：{dst.name} -> {src.parent}')
        # 清理空文件夹
        for folder in archive_root.iterdir():
            if folder.is_dir() and folder.name != 'undo_log.txt' and not any(folder.iterdir()):
                folder.rmdir()
                self.log(f'已删除空文件夹：{folder.name}')
        log_path.unlink()
        self.log("撤回完成，日志已删除。\n")
        messagebox.showinfo("完成", "撤回成功！")
        self.refresh_list()


# ---------- 程序入口 ----------
if __name__ == '__main__':
    root = tk.Tk()
    app = DeskOrganizerApp(root)
    root.mainloop()